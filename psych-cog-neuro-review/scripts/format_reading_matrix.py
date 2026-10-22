"""把阅读矩阵整理成给人看的 Excel。

数据来源是 CSV（如 文献阅读矩阵_literature_matrix.csv），或 xlsx 的“数据”工作表。
输出 xlsx 里有两个给人看的工作表：
- 阅读：前几列先给人看（文献、一句话结论、研究概况、主要章节、证据强度等），
  其余编码字段按组排在后面并默认折叠；等级代码自动带上中文说明；
  按文章结构分组，每组前加小节标题行。点“文献”打开阅读笔记，点“全文”打开 PDF。
  格子里的双链（如 [[正文.pdf#page=3|第 3 页]]）只显示文字“第 3 页”，因为 Excel 打不开双链。
- 说明：每个字段的中文名和分组，以及等级代码的含义。
输入是 xlsx 时，原来的数据留在“数据”工作表，脚本不改动它；输入是 CSV 时，数据仍以 CSV 为准。

以后要改分组、中文名或列宽，改本文件顶部的 GROUPS、LABELS 和 WIDE 三处。

用法：
  python format_reading_matrix.py --矩阵 <csv或xlsx> [--输出 <xlsx>] [--知识库 <知识库文件夹>]
      [--文章结构 <md>] [--等级说明 <md> ...] [--全文目录 <PDF文件夹>]
"""
from __future__ import annotations

import argparse
import csv
import re
import urllib.parse
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 给人看的前几列：显示名 -> 依次查找的数据字段
HUMAN = [
    ("文献", ["citation", "文献"]),
    ("一句话结论", ["one_sentence_conclusion", "一句话结论"]),
    ("研究概况", ["study_snapshot", "研究概况"]),
    ("主要章节", ["primary_section", "主要章节"]),
    ("其他章节与主张", ["other_sections_and_claims", "其他章节与主张"]),
    ("证据强度", ["evidence_strength", "证据强度"]),
    ("阅读深度", ["reading_level", "阅读深度"]),
    ("状态", ["提取状态", "核对状态", "professor_briefing_status", "reviewer_b_status"]),
    ("全文", ["source_pdf", "全文"]),
]
HUMAN_FIELDS = {name for _, names in HUMAN for name in names} | {"note_path", "笔记"}

# 编码字段分组：组名、表头底色、字段（英文键名和中文字段名都可以）
GROUPS = [
    ("身份", "DDEBF7", ["study_id", "candidate_id", "report_id", "study_family_id", "doi", "zotero_item_key",
                       "title", "authors", "year", "候选编号", "研究编号", "报告编号", "研究家族编号",
                       "第一作者", "年份", "题名", "期刊", "DOI"]),
    ("样本", "E2EFDA", ["country", "sample", "sample_n", "age_or_grade", "国家或地区", "人群年龄段", "年龄",
                       "样本类型", "诊断类别", "样本量", "合并或共享样本"]),
    ("测量", "FFF2CC", ["constructs", "measures", "measurement_reliability_validity", "scoring_details",
                       "测量工具", "测量体系", "C等级", "信息源", "指导语时间框架", "领域分及其可靠性"]),
    ("结局", "FCE4D6", ["结局域", "结局与工具", "结局信息源", "与人格功能条目重叠", "是否人格障碍特征"]),
    ("设计与分析", "EDE2F6", ["design", "waves", "wave_intervals", "intervals", "analysis", "analysis_model",
                          "estimand_level", "covariates", "missing_data_handling", "data_source",
                          "data_cleaning_exclusions", "model_diagnostics", "robustness_sensitivity",
                          "研究设计", "波数与间隔", "分析层次", "分析模型", "控制变量", "控制方式"]),
    ("结果", "D9E1F2", ["main_findings", "primary_results", "null_opposite_mixed_results", "effect_data",
                       "effect_estimates", "fulltext_location", "evidence_location",
                       "主要结果", "效应量", "结果方向", "原文摘录", "证据位置"]),
    ("分级与用途", "FFE699", ["claim_ceiling", "construct_match", "E等级", "M等级", "新问题编号", "review_role",
                          "paragraph_level_use", "relation_to_project_results", "research_question",
                          "study_hypotheses", "claimed_novelty", "theory_framework", "author_claim",
                          "results_demonstrate", "reader_inference", "key_references_to_trace",
                          "briefing_2min", "anticipated_professor_questions"]),
    ("同分异质与发展", "E2EFDA", ["同分设计类型", "独立效标类型", "年龄效应", "测量等值", "常模参照", "是否测成熟度"]),
    ("质量与偏倚", "F8CBAD", ["quality", "limitations", "bias_selection", "bias_measurement", "bias_temporality",
                          "bias_confounding", "bias_missingness", "bias_analysis", "bias_reporting",
                          "applicability", "极端组", "预注册", "主要局限"]),
    ("阅读过程", "EDEDED", ["reading_scope", "pass1_orientation_status", "pass2_methods_results_status",
                         "pass3_critical_writing_status", "legacy_reading_status", "abstract_status",
                         "introduction_status", "methods_status", "data_status", "results_status",
                         "discussion_status", "conclusion_status", "references_trace_status",
                         "abstract_audit", "introduction_gap", "procedure_ethics", "discussion_interpretation",
                         "paper_type", "reporting_standard_applicable", "source_package_status",
                         "consistency_audit_status", "key_figure_ids", "figure_claim_audit",
                         "figure_data_traceability", "figure_redraw_or_reuse_plan", "family_status",
                         "reviewer_a_status", "reviewer_b_status", "adjudication_status", "screening_status",
                         "verified_at", "证据来源级别", "提取人", "提取日期", "核对日期"]),
]
OTHER_GROUP = ("其他", "F2F2F2")
V1_GROUP = ("v1 参考", "D9D9D9")

# 英文键名的中文名；中文字段名直接显示
LABELS = {
    "study_id": "研究编号", "candidate_id": "候选编号", "report_id": "报告编号", "study_family_id": "研究家族编号",
    "citation": "文献", "doi": "DOI", "zotero_item_key": "Zotero 条目", "title": "题名", "authors": "作者",
    "year": "年份", "country": "国家或地区", "sample": "样本", "sample_n": "样本量", "age_or_grade": "年龄或年级",
    "constructs": "构念", "measures": "测量工具", "measurement_reliability_validity": "信效度",
    "scoring_details": "计分", "design": "设计", "waves": "波数", "wave_intervals": "波次间隔",
    "intervals": "波次间隔", "analysis": "分析", "analysis_model": "分析模型", "estimand_level": "估计层次",
    "covariates": "协变量", "missing_data_handling": "缺失处理", "data_source": "数据来源",
    "data_cleaning_exclusions": "清理与排除", "model_diagnostics": "模型诊断",
    "robustness_sensitivity": "稳健性与敏感性", "main_findings": "主要结果", "primary_results": "主要结果",
    "null_opposite_mixed_results": "零、相反和混合结果", "effect_data": "效应数据",
    "effect_estimates": "效应估计", "fulltext_location": "证据位置", "evidence_location": "证据位置",
    "claim_ceiling": "主张上限", "construct_match": "构念匹配", "review_role": "综述角色",
    "paragraph_level_use": "段落用途", "relation_to_project_results": "与本研究结果的关系",
    "research_question": "研究问题", "study_hypotheses": "研究假设", "claimed_novelty": "作者声称的创新",
    "theory_framework": "理论框架", "author_claim": "作者主张", "results_demonstrate": "结果实际显示",
    "reader_inference": "我的推断", "key_references_to_trace": "待追踪文献", "briefing_2min": "两分钟汇报",
    "anticipated_professor_questions": "可能的追问", "quality": "质量", "limitations": "局限",
    "bias_selection": "偏倚：抽样", "bias_measurement": "偏倚：测量", "bias_temporality": "偏倚：时间顺序",
    "bias_confounding": "偏倚：混杂", "bias_missingness": "偏倚：缺失", "bias_analysis": "偏倚：分析",
    "bias_reporting": "偏倚：选择性报告", "applicability": "适用性", "reading_scope": "阅读范围",
    "pass1_orientation_status": "第一轮状态", "pass2_methods_results_status": "第二轮状态",
    "pass3_critical_writing_status": "第三轮状态", "legacy_reading_status": "旧阅读状态",
    "abstract_status": "摘要状态", "introduction_status": "引言状态", "methods_status": "方法状态",
    "data_status": "数据状态", "results_status": "结果状态", "discussion_status": "讨论状态",
    "conclusion_status": "结论状态", "references_trace_status": "参考文献追踪状态",
    "professor_briefing_status": "汇报准备状态", "abstract_audit": "摘要反核", "introduction_gap": "研究缺口",
    "procedure_ethics": "程序与伦理", "discussion_interpretation": "讨论与解释", "paper_type": "论文类型",
    "reporting_standard_applicable": "适用报告规范", "source_package_status": "全文包状态",
    "consistency_audit_status": "一致性核查", "key_figure_ids": "关键图", "figure_claim_audit": "图与主张核查",
    "figure_data_traceability": "图数据可追溯性", "figure_redraw_or_reuse_plan": "重绘或复用计划",
    "family_status": "家族状态", "reviewer_a_status": "第一人状态", "reviewer_b_status": "第二人状态",
    "adjudication_status": "裁决状态", "screening_status": "筛选状态", "verified_at": "核对日期",
    "source_pdf": "全文", "reading_level": "阅读层级", "one_sentence_conclusion": "一句话结论",
    "study_snapshot": "研究概况", "primary_section": "主要章节", "other_sections_and_claims": "其他章节与主张",
    "evidence_strength": "证据强度", "note_path": "笔记", "notes": "备注",
}
WIDE = {"一句话结论": 60, "研究概况": 40, "证据强度": 36, "其他章节与主张": 24, "主要结果": 44,
        "原文摘录": 44, "主要局限": 36, "零、相反和混合结果": 40, "效应估计": 30, "效应量": 30}
HUMAN_WIDTH = {"文献": 18, "主要章节": 16, "阅读深度": 9, "状态": 18, "全文": 8}
STATUS_FILL = [("受阻", "F8CBAD"), ("冲突", "F8CBAD"), ("待核", "FFF2CC"), ("待第二人", "FFF2CC"),
               ("已核实", "C6EFCE"), ("核对通过", "C6EFCE"), ("就绪", "C6EFCE")]


WIKILINK = re.compile(r"!?\[\[([^\]|]+?)(?:\|([^\]]+))?\]\]")


def plain(text: str) -> str:
    """把双链换成显示文字：[[路径#page=3|第 3 页]] → 第 3 页；没有显示文字的取文件名。"""
    def repl(match: re.Match) -> str:
        target, alias = match.group(1), match.group(2)
        if alias:
            return alias
        name, _, anchor = target.partition("#")
        name = Path(name).stem if name else ""
        page = re.match(r"page=(\d+)", anchor)
        return f"{name} 第 {page.group(1)} 页" if page else (name or anchor)
    return WIKILINK.sub(repl, text)


def read_rows(path: Path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
        return None, None, rows[0], [r + [""] * (len(rows[0]) - len(r)) for r in rows[1:]]
    wb = load_workbook(path)
    data_ws = wb["数据"] if "数据" in wb.sheetnames else next(
        wb[name] for name in wb.sheetnames if name not in ("阅读", "说明"))
    values = list(data_ws.iter_rows(values_only=True))
    header = ["" if v is None else str(v) for v in values[0]]
    body = [["" if v is None else str(v) for v in row] for row in values[1:] if any(v not in (None, "") for v in row)]
    return wb, data_ws, header, body


def code_labels(paths: list[Path]) -> dict[str, str]:
    labels: dict[str, str] = {}
    for path in paths:
        for line in path.read_text(encoding="utf-8").splitlines():
            cells = [c.strip() for c in line.strip().strip("|").split("|")] if line.startswith("|") else []
            if len(cells) < 2:
                continue
            first = re.fullmatch(r"([A-Z]\d{1,2})(?:\s+(.+))?", cells[0])
            if not first:
                continue
            meaning = first.group(2) or cells[1]
            meaning = re.sub(r"[*`]", "", meaning).strip()
            if meaning and not meaning.startswith("-"):
                labels.setdefault(first.group(1), meaning)
    return labels


def section_order(path: Path | None) -> list[tuple[str, str]]:
    if not path:
        return []
    order = []
    lines = path.read_text(encoding="utf-8").splitlines()
    if "## 结构" in (line.strip() for line in lines):  # 跳过变更记录等前文
        lines = lines[[line.strip() for line in lines].index("## 结构"):]
    for line in lines:
        chapter = re.match(r"^\*\*(\d+)\s+([^*]+)\*\*", line.strip())
        sub = re.match(r"^\s*-\s+(\d+\.\d+)\s+(.+)$", line)
        if chapter:
            order.append((chapter.group(1), re.split(r"[（(]", chapter.group(2))[0].strip()))
        elif sub:
            order.append((sub.group(1), re.split(r"[：:（(]", sub.group(2))[0].strip()))
    seen, unique = set(), []
    for num, title in order:
        if num not in seen:
            seen.add(num)
            unique.append((num, title))
    return unique


def find_notes(kb: Path | None) -> dict[str, Path]:
    notes: dict[str, Path] = {}
    if not kb or not (kb / "文献").is_dir():
        return notes
    for path in (kb / "文献").rglob("*.md"):
        head = path.read_text(encoding="utf-8").split("\n---", 1)[0]
        match = re.search(r"^候选编号:\s*(.+)$", head, flags=re.M)
        if match:
            notes[match.group(1).strip().strip("\"'")] = path
    return notes


def note_summary(path: Path) -> str:
    match = re.search(r"\*\*一句话结论\*\*[：:]\s*(.+)", path.read_text(encoding="utf-8"))
    return match.group(1).strip() if match else ""


def obsidian_uri(path: Path) -> str | None:
    for parent in path.parents:
        if (parent / ".obsidian").is_dir():
            rel = path.relative_to(parent).with_suffix("").as_posix()
            query = urllib.parse.urlencode({"vault": parent.name, "file": rel}, quote_via=urllib.parse.quote)
            return f"obsidian://open?{query}"
    return None


def main() -> None:
    parser = argparse.ArgumentParser(description="整理阅读矩阵的阅读页")
    parser.add_argument("--矩阵", "--matrix", dest="matrix", required=True)
    parser.add_argument("--输出", "--output", dest="output")
    parser.add_argument("--知识库", "--vault", dest="kb")
    parser.add_argument("--文章结构", "--outline", dest="outline")
    parser.add_argument("--等级说明", "--codes", dest="codes", action="append", default=[])
    parser.add_argument("--全文目录", "--pdf-dir", dest="pdf_dir")
    args = parser.parse_args()

    source = Path(args.matrix).resolve()
    wb, data_ws, header, rows = read_rows(source)
    output = Path(args.output).resolve() if args.output else source.with_suffix(".xlsx")
    col = {name: i for i, name in enumerate(header)}
    labels = code_labels([Path(p) for p in args.codes])
    order = section_order(Path(args.outline) if args.outline else None)
    notes = find_notes(Path(args.kb).resolve() if args.kb else None)
    pdfs: dict[str, Path] = {}
    if args.pdf_dir and Path(args.pdf_dir).is_dir():
        for pdf in sorted(Path(args.pdf_dir).glob("*.pdf")):
            pdfs.setdefault(pdf.stem.split("_")[0], pdf)

    def get(row, names):
        for name in names:
            if name in col and row[col[name]].strip():
                return row[col[name]].strip()
        return ""

    def with_label(value: str) -> str:
        parts = [p.strip() for p in re.split(r"[；;]", value) if p.strip()]
        if parts and all(p in labels for p in parts):
            return "；".join(f"{p} {labels[p]}" for p in parts)
        return value

    # 编码字段分组
    grouped: list[tuple[str, str, list[str]]] = []
    used = set(HUMAN_FIELDS)
    for name, color, fields in GROUPS:
        present = [f for f in header if f in fields and f not in used]
        used.update(present)
        if present:
            grouped.append((name, color, present))
    v1 = [f for f in header if f.startswith("v1") and f not in used]
    used.update(v1)
    rest = [f for f in header if f not in used and f]
    if rest:
        grouped.append((OTHER_GROUP[0], OTHER_GROUP[1], rest))
    if v1:
        grouped.append((V1_GROUP[0], V1_GROUP[1], v1))

    # 排序：先按文章结构，再按阅读深度、E 等级、年份
    rank_depth = {"D1": 0, "4级": 0, "D2": 1, "3级": 1, "D3": 2, "2级": 2}
    known = [num for num, _ in order]

    def section_of(row) -> str:
        match = re.search(r"\d+(?:\.\d+)?", get(row, ["primary_section", "主要章节"]))
        return match.group(0) if match else ""

    def sort_key(row):
        sec = section_of(row)
        pos = known.index(sec) if sec in known else (len(known) + 1 if not sec else len(known))
        e_level = max([int(x) for x in re.findall(r"E(\d)", get(row, ["E等级"]))] or [0])
        year = re.search(r"\d{4}", get(row, ["年份", "year"]) or get(row, ["citation", "文献"]))
        return (pos, sec, rank_depth.get(get(row, ["reading_level", "阅读深度"]), 3), -e_level,
                -(int(year.group(0)) if year else 0))

    rows.sort(key=sort_key)

    out_wb = wb if wb is not None else Workbook()
    if wb is not None:
        if data_ws.title != "数据":
            data_ws.title = "数据"
        for name in ("阅读", "说明"):
            if name in out_wb.sheetnames:
                del out_wb[name]
        ws = out_wb.create_sheet("阅读", 0)
    else:
        ws = out_wb.active
        ws.title = "阅读"

    human_names = [name for name, _ in HUMAN]
    columns = [(name, None) for name in human_names]
    for _, _, fields in grouped:
        columns += [(LABELS.get(f, f), f) for f in fields]

    band = PatternFill("solid", fgColor="BDD7EE")
    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    ws.cell(1, 1, "先看这几列").font = bold
    ws.cell(1, 1).fill = band
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(human_names))
    start = len(human_names) + 1
    spans = []
    for name, color, fields in grouped:
        end = start + len(fields) - 1
        cell = ws.cell(1, start, name)
        cell.font = bold
        cell.fill = PatternFill("solid", fgColor=color)
        if end > start:
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        spans.append((start, end, color))
        start = end + 1
    for j, (title, _) in enumerate(columns, 1):
        cell = ws.cell(2, j, title)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = band if j <= len(human_names) else next(
            PatternFill("solid", fgColor=c) for s, e, c in spans if s <= j <= e)

    titles = dict(order)
    r = 3
    current = None
    section_fill = PatternFill("solid", fgColor="1F4E78")
    for row in rows:
        sec = section_of(row)
        if sec != current:
            current = sec
            count = sum(1 for x in rows if section_of(x) == sec)
            text = f"{sec} {titles.get(sec, '')}（{count} 篇）".strip() if sec else f"尚未定位章节（{count} 篇）"
            cell = ws.cell(r, 1, text)
            cell.font = Font(bold=True, color="FFFFFF")
            for j in range(1, len(columns) + 1):
                ws.cell(r, j).fill = section_fill
            r += 1
        cid = get(row, ["candidate_id", "候选编号"])
        note = notes.get(cid)
        author = get(row, ["第一作者"]).split(",")[0].strip()
        year = get(row, ["年份", "year"])
        cite = get(row, ["citation", "文献"]) or (f"{author} {year}".strip() if author else cid)
        values = {
            "文献": cite,
            "一句话结论": get(row, ["one_sentence_conclusion", "一句话结论"]) or (note_summary(note) if note else ""),
            "研究概况": get(row, ["study_snapshot", "研究概况"]),
            "主要章节": get(row, ["primary_section", "主要章节"]),
            "其他章节与主张": get(row, ["other_sections_and_claims", "其他章节与主张"]),
            "证据强度": get(row, ["evidence_strength", "证据强度"]),
            "阅读深度": get(row, ["reading_level", "阅读深度"]),
            "状态": "｜".join(v for v in (row[col[n]].strip() for n in HUMAN[7][1] if n in col) if v),
            "全文": "打开" if (cid in pdfs or get(row, ["source_pdf", "全文"])) else "",
        }
        for j, (title, field) in enumerate(columns, 1):
            value = values[title] if field is None else with_label(row[col[field]].strip())
            cell = ws.cell(r, j, plain(value))
            cell.alignment = wrap
        if note and (uri := obsidian_uri(note)):
            ws.cell(r, 1).hyperlink = uri
            ws.cell(r, 1).font = Font(color="0563C1", underline="single")
        pdf = pdfs.get(cid) or (Path(get(row, ["source_pdf", "全文"])) if get(row, ["source_pdf", "全文"]) else None)
        if pdf:
            ws.cell(r, 9).hyperlink = pdf.resolve().as_uri() if pdf.is_absolute() or pdf.exists() else str(pdf)
            ws.cell(r, 9).font = Font(color="0563C1", underline="single")
        status = values["状态"]
        for key, color in STATUS_FILL:
            if key in status:
                ws.cell(r, 8).fill = PatternFill("solid", fgColor=color)
                break
        r += 1

    for j, (title, _) in enumerate(columns, 1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = HUMAN_WIDTH.get(title) or WIDE.get(title, 18)
    for j in range(len(human_names) + 1, len(columns) + 1):  # 编码字段统一折叠，点 + 展开
        dim = ws.column_dimensions[get_column_letter(j)]
        dim.outline_level = 1
        dim.hidden = True
    ws.freeze_panes = "C3"
    ws.sheet_view.zoomScale = 100

    notes_ws = out_wb.create_sheet("说明", 1)
    notes_ws.append(["怎么看"])
    notes_ws.append(["“阅读”页的前几列给人看；后面是按组排列的编码字段，默认折叠，点表格上方的 + 全部展开，第 1 行写着每组的组名。"])
    notes_ws.append(["行按文章结构分组；点“文献”打开阅读笔记，点“全文”打开 PDF。本页和“阅读”页由脚本生成，改数据请改“数据”页或原 CSV。"])
    notes_ws.append([])
    notes_ws.append(["分组", "字段", "数据里的列名"])
    for name, names in HUMAN:
        found = [n for n in names if n in col]
        notes_ws.append(["先看这几列", name, "、".join(found) or "（脚本生成）"])
    for name, _, fields in grouped:
        for f in fields:
            notes_ws.append([name, LABELS.get(f, f), f])
    if labels:
        notes_ws.append([])
        notes_ws.append(["代码", "含义"])
        for key in sorted(labels):
            notes_ws.append([key, labels[key]])
    for letter, width in (("A", 16), ("B", 60), ("C", 36)):
        notes_ws.column_dimensions[letter].width = width
    for row in notes_ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    out_wb.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(output)
    print(f"已生成：{output}（{len(rows)} 篇，{len(columns)} 列，其中给人看的 {len(human_names)} 列）")


if __name__ == "__main__":
    main()
