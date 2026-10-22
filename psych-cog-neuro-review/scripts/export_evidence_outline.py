"""从证据提纲（每章一个 Markdown）导出给人看的 Excel 和机器用的主张证据对应表。

证据提纲的写法见 templates/ 里的证据提纲模板，要点：
  ## 4.3 小节标题
  ### 第 2 段　段落要点
  段落作用：……
  #### 主张 4.3-2-1
  拟写进正文的一句话
  - **级别与上限**：……
  - **证据**：
    - 支持｜[[笔记名#^f2|作者 年份 发现 2]]：一句话说明（第 7 页，表 2）
  - **证据强度**：……  - **可以写成**：……  - **不能写成**：……  - **缺口**：……  - **状态**：……

证据方向只用：支持、零结果、相反、混合、限定、背景（实证论文的结果主张另可用“本研究结果”）。
提纲是唯一来源；两个导出文件每次都整体重写，不要手改。

导出时顺带自检，只提示、不改提纲：证据行有没有链到笔记里的发现块、块编号在不在、行末有没有页码；
已挂证据的主张，状态是否还停在“待填证据”，证据概述、可以写成、不能写成、缺口是否还空着。
发现的问题写进总表的“待修”页，并在命令行列出。

用法：
  python export_evidence_outline.py --提纲目录 <文件夹> [--知识库 <知识库文件夹>]
      [--输出表 <xlsx>] [--输出csv <主张证据对应表_claim_evidence_map.csv>]
"""
from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CLAIM_HEADER = [
    "claim_id", "claim_text", "claim_type", "candidate_id", "study_id", "report_id", "study_family_id",
    "support_status", "claim_ceiling", "support_carrier_type", "support_carrier_value", "fulltext_location",
    "construct_match", "estimand_level", "result_direction", "reviewer_status", "reviewer_reason", "verified_at",
    "effect_estimate", "standard_error", "ci_95", "p_value", "paper_one_sentence_conclusion",
    "paragraph_function", "cautious_paraphrase", "comparison_or_companion_sources",
    "relation_to_project_results", "prohibited_wording", "manuscript_destination", "paragraph_id",
    "argument_role", "warrant_text", "counterevidence_handling", "drafting_status", "figure_ids", "figure_role",
]
DIRECTION = {"支持": "supports", "零结果": "null", "相反": "opposite", "反对": "opposite", "混合": "mixed",
             "限定": "mixed", "背景": "not-tested", "本研究结果": "supports"}
CEILING = [("因果", "causal-effect"), ("机制", "mechanism-candidate"), ("个体内", "within-person-temporal"),
           ("时间预测", "temporal-prediction"), ("前瞻", "temporal-prediction"), ("关联", "association"),
           ("描述", "descriptive"), ("定义", "definition")]
CLAIM_TYPE = {"定义": "definition", "理论": "theory", "患病率": "prevalence", "直接实证": "direct-empirical",
              "测量": "measurement", "方法": "method", "背景": "context", "解释": "interpretation"}
FIELD = re.compile(r"^- \*\*(.+?)\*\*[：:]\s*(.*)$")
EVIDENCE = re.compile(r"^\s+-\s+(支持|零结果|相反|反对|混合|限定|背景|本研究结果)\s*｜\s*(.+)$")
LINK = re.compile(r"\[\[([^\]|#]+)(?:#([^\]|]*))?(?:\|([^\]]*))?\]\]")
LOCATION = re.compile(r"[（(]([^（）()]*(?:页|表|图)[^（）()]*)[）)]\s*[。.；;]?\s*$")


def frontmatter(text: str) -> dict[str, str]:
    if not text.startswith("---"):
        return {}
    head = text.split("\n---", 1)[0]
    return {m.group(1): m.group(2).strip().strip("\"'") for m in re.finditer(r"^([^:\n]+):\s*(.*)$", head, re.M)}


def load_notes(kb: Path | None) -> dict[str, dict[str, str]]:
    notes: dict[str, dict[str, str]] = {}
    if not kb or not (kb / "文献").is_dir():
        return notes
    for path in (kb / "文献").rglob("*.md"):
        text = path.read_text(encoding="utf-8")
        meta = frontmatter(text)
        summary = re.search(r"\*\*一句话结论\*\*[：:]\s*(.+)", text)
        meta["一句话结论"] = summary.group(1).strip() if summary else ""
        meta["块编号"] = " ".join(sorted(set(re.findall(r"\^([\w-]+)\s*$", text, re.M))))
        notes[path.stem] = meta
    return notes


def check(claims: list[dict], notes: dict[str, dict[str, str]], have_kb: bool) -> list[tuple[str, str, str]]:
    """找出常见的挂接问题，返回（主张编号，问题，涉及的内容）。只提示，不改提纲。"""
    issues = []
    for item in claims:
        f = item["fields"]
        if item["evidence"]:
            if "待填证据" in f.get("状态", ""):
                issues.append((item["id"], "已挂证据，状态还是“待填证据”", ""))
            empty = [k for k in ("证据概述", "证据强度", "可以写成", "不能写成", "缺口") if not f.get(k)]
            if empty:
                issues.append((item["id"], "已挂证据，这几栏还空着：" + "、".join(empty), ""))
        for _, text in item["evidence"]:
            link = LINK.search(text)
            if not link:
                issues.append((item["id"], "证据行没有链接阅读笔记", text[:60]))
                continue
            target, anchor = link.group(1), link.group(2) or ""
            if not anchor.startswith("^"):
                issues.append((item["id"], "没有链到笔记里的发现块（如 #^f2）", text[:60]))
            if have_kb:
                if target not in notes:
                    issues.append((item["id"], "链接的阅读笔记不存在", target))
                elif anchor.startswith("^") and anchor[1:] not in notes[target].get("块编号", "").split():
                    issues.append((item["id"], f"笔记里找不到块编号 {anchor}", target))
            if not LOCATION.search(text):
                issues.append((item["id"], "行末括号里没有页码、表或图", text[:60]))
    return issues


def parse(path: Path) -> list[dict]:
    claims: list[dict] = []
    section = paragraph = ""
    section_title = paragraph_title = paragraph_role = ""
    claim = None
    field = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if m := re.match(r"^##\s+(\d+(?:\.\d+)*)\s+(.*)$", line):
            section, section_title, paragraph, claim, field = m.group(1), m.group(2).strip(), "", None, None
            continue
        if m := re.match(r"^###\s+第\s*(\d+)\s*段\s*(.*)$", line):
            paragraph, paragraph_title, paragraph_role, claim, field = m.group(1), m.group(2).strip(), "", None, None
            continue
        if m := re.match(r"^段落作用[：:]\s*(.*)$", line):
            paragraph_role = m.group(1).strip()
            continue
        if m := re.match(r"^####\s+主张\s+(\S+)\s*$", line):
            claim = {"id": m.group(1), "text": "", "section": section, "section_title": section_title,
                     "paragraph": paragraph, "paragraph_title": paragraph_title, "role": paragraph_role,
                     "fields": {}, "evidence": [], "file": path.name}
            claims.append(claim)
            field = None
            continue
        if claim is None or not line.strip():
            continue
        if m := FIELD.match(line):
            field = m.group(1).strip()
            claim["fields"][field] = m.group(2).strip()
            continue
        if m := EVIDENCE.match(line):
            claim["evidence"].append((m.group(1), m.group(2).strip()))
            continue
        if field is None and not line.lstrip().startswith(("-", "#", ">")):
            claim["text"] = (claim["text"] + line.strip()).strip()
        elif field is not None and line.startswith(("  ", "\t")):
            claim["fields"][field] = (claim["fields"][field] + " " + line.strip()).strip()
    return claims


def main() -> None:
    parser = argparse.ArgumentParser(description="导出证据提纲")
    parser.add_argument("--提纲目录", "--outline-dir", dest="folder", required=True)
    parser.add_argument("--知识库", "--vault", dest="kb")
    parser.add_argument("--输出表", "--xlsx", dest="xlsx")
    parser.add_argument("--输出csv", "--csv", dest="csv")
    args = parser.parse_args()

    folder = Path(args.folder).resolve()
    files = []
    for path in folder.glob("*.md"):
        meta = frontmatter(path.read_text(encoding="utf-8"))
        if meta.get("类型") == "证据提纲" or "证据提纲" in path.stem:
            chapter = re.search(r"\d+", meta.get("章", "") or path.stem)
            files.append((int(chapter.group(0)) if chapter else 999, path.name, path, meta))
    files.sort()
    notes = load_notes(Path(args.kb).resolve() if args.kb else None)

    claims, chapters = [], []
    for number, _, path, meta in files:
        chapters.append((number, meta.get("标题", path.stem)))
        for item in parse(path):
            item["chapter"] = number
            item["chapter_title"] = meta.get("标题", path.stem)
            claims.append(item)

    xlsx = Path(args.xlsx).resolve() if args.xlsx else folder / "证据提纲总表.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "证据提纲"
    headers = ["段落", "主张编号", "主张", "级别与上限", "证据概述", "证据", "证据强度",
               "可以写成", "不能写成", "缺口", "状态"]
    widths = [26, 11, 40, 20, 40, 60, 28, 40, 30, 30, 14]
    ws.append(headers)
    for j, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = width
        ws.cell(1, j).font = Font(bold=True)
        ws.cell(1, j).fill = PatternFill("solid", fgColor="BDD7EE")
    wrap = Alignment(wrap_text=True, vertical="top")
    chapter_fill, section_fill = PatternFill("solid", fgColor="1F4E78"), PatternFill("solid", fgColor="DDEBF7")
    last_chapter = last_section = None
    for item in claims:
        if item["chapter"] != last_chapter:
            last_chapter, last_section = item["chapter"], None
            ws.append([f"第 {item['chapter']} 章　{item['chapter_title']}"])
            for j in range(1, len(headers) + 1):
                ws.cell(ws.max_row, j).fill = chapter_fill
            ws.cell(ws.max_row, 1).font = Font(bold=True, color="FFFFFF")
        if item["section"] != last_section:
            last_section = item["section"]
            ws.append([f"{item['section']} {item['section_title']}"])
            for j in range(1, len(headers) + 1):
                ws.cell(ws.max_row, j).fill = section_fill
            ws.cell(ws.max_row, 1).font = Font(bold=True)
        f = item["fields"]
        evidence = "\n".join(f"{d}｜{LINK.sub(lambda m: m.group(3) or m.group(1), text)}" for d, text in item["evidence"])
        para = f"第 {item['paragraph']} 段　{item['paragraph_title']}" + (f"\n（{item['role']}）" if item["role"] else "")
        ws.append([para, item["id"], item["text"], f.get("级别与上限", ""), f.get("证据概述", ""), evidence,
                   f.get("证据强度", ""), f.get("可以写成", ""), f.get("不能写成", ""), f.get("缺口", ""),
                   f.get("状态", "")])
        for j in range(1, len(headers) + 1):
            ws.cell(ws.max_row, j).alignment = wrap
    ws.freeze_panes = "C2"

    by_paper = wb.create_sheet("文献到主张")
    by_paper.append(["文献", "候选编号", "主张编号", "方向", "说明"])
    pairs = []
    for item in claims:
        for direction, text in item["evidence"]:
            for target, _, _ in LINK.findall(text)[:1]:
                pairs.append((target, notes.get(target, {}).get("候选编号", ""), item["id"], direction,
                              LINK.sub(lambda m: m.group(3) or m.group(1), text)))
    for row in sorted(pairs):
        by_paper.append(list(row))
    for j, width in enumerate([34, 18, 12, 8, 80], 1):
        by_paper.column_dimensions[get_column_letter(j)].width = width
        by_paper.cell(1, j).font = Font(bold=True)
    for row in by_paper.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap

    issues = check(claims, notes, bool(args.kb))
    if issues:
        todo = wb.create_sheet("待修")
        todo.append(["主张编号", "问题", "涉及"])
        for row in issues:
            todo.append(list(row))
        for j, width in enumerate([12, 44, 60], 1):
            todo.column_dimensions[get_column_letter(j)].width = width
            todo.cell(1, j).font = Font(bold=True)
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xlsx)

    if args.csv:
        out = Path(args.csv).resolve()
        rows = []
        for item in claims:
            f = item["fields"]
            status = f.get("状态", "")
            ceiling = next((v for k, v in CEILING if k in f.get("级别与上限", "")), "")
            reviewer = ("conflict" if "冲突" in status else "adjudicated" if "裁决" in status
                        else "agreed" if ("一致" in status or "核对通过" in status) else "A-only/B-pending")
            counter = "；".join(LINK.sub(lambda m: m.group(3) or m.group(1), t)
                               for d, t in item["evidence"] if d in ("相反", "反对", "零结果", "限定"))
            base = {
                "claim_id": item["id"], "claim_text": item["text"],
                "claim_type": CLAIM_TYPE.get(f.get("主张类型", ""), ""), "claim_ceiling": ceiling,
                "reviewer_status": reviewer, "paragraph_function": item["role"],
                "cautious_paraphrase": f.get("可以写成", ""), "prohibited_wording": f.get("不能写成", ""),
                "manuscript_destination": f"{item['section']} {item['section_title']}",
                "paragraph_id": f"{item['section']}-{item['paragraph']}", "drafting_status": status,
                "counterevidence_handling": counter or f.get("反证与限定", ""),
                "warrant_text": f.get("证据概述", ""),
            }
            if not item["evidence"]:
                rows.append({**base, "support_status": "blocked", "result_direction": "not-tested"})
                continue
            for direction, text in item["evidence"]:
                target = (LINK.findall(text) or [("", "", "")])[0][0]
                meta = notes.get(target, {})
                where = LOCATION.search(text)
                rows.append({
                    **base,
                    "candidate_id": meta.get("候选编号", ""), "study_id": meta.get("研究编号", ""),
                    "study_family_id": meta.get("研究家族编号", ""),
                    "support_status": ("confirmed" if "已核实" in status and direction == "支持"
                                       else "rejected" if direction in ("相反", "反对") else "partial"),
                    "result_direction": DIRECTION[direction],
                    "fulltext_location": where.group(1) if where else "",
                    "support_carrier_type": "multiple" if where and "表" in where.group(1) and "页" in where.group(1) else "",
                    "paper_one_sentence_conclusion": meta.get("一句话结论", ""),
                    "relation_to_project_results": LINK.sub(lambda m: m.group(3) or m.group(1), text)
                    if direction == "本研究结果" else "",
                })
        out.parent.mkdir(parents=True, exist_ok=True)
        with out.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=CLAIM_HEADER, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({key: row.get(key, "") for key in CLAIM_HEADER})
        print(f"主张证据对应表：{out}（{len(rows)} 行）")
    print(f"证据提纲总表：{xlsx}（{len(chapters)} 章，{len(claims)} 条主张）")
    if issues:
        print(f"提纲自检：{len(issues)} 处待修，已写进总表的“待修”页。前几处：")
        for claim_id, problem, detail in issues[:8]:
            print(f"  主张 {claim_id}：{problem}" + (f"｜{detail}" if detail else ""))
    elif args.kb:
        print("提纲自检：没有发现问题")


if __name__ == "__main__":
    main()
