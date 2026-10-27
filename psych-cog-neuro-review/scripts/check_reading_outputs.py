#!/usr/bin/env python3
"""阅读笔记、知识卡和阅读矩阵的自检。交稿前运行，报告里本篇的问题要清零。

查五类问题，结果写进一个 xlsx（汇总、笔记、知识卡、数字核对、坏链接、矩阵），命令行打印各类数量：
1. 版式：速览框是否在开头、该有的小节是否齐、篇幅够不够、工具和方法有没有补课框、标题是否写成了限制；
2. 讲法：句号后或段尾连着的页码、一段里多处限定、讲解区出现项目代号或内部用语（含“待核”“不擅自”这类核对用语）、
   用中文数字计数、块编号前缺空格、收集限定句的小节（如“方法与解释补充”）、教学示例标签不是“教学示例（虚构）”、
   公式写成 \\( \\) 或 \\[ \\]（Obsidian 不显示）、控制字符（脚本把 \\f 这类转义吃掉了）；知识卡里，完整的方法讲义要有公式；
3. 链接：双链、嵌入的小节、块编号、PDF 页码、图片是否都找得到；
4. 数字：段落里的统计数字能不能在它所引的 PDF 页（前后各放宽 1 页）上找到。找不到的不一定错：
   可能是笔记换算过（应写成“约……”或写明怎么算），可能来自图（文字层里没有），可能页码标错了；
5. 矩阵（加 --矩阵）：受控取值、计数用的名称是否对得上知识库节点、占位词、双链、与笔记是否一致。
   受控字段和取值从项目的取值表读，格式见 references/matrix-vocabulary-template.json。

用法：
  python check_reading_outputs.py --根 <Obsidian 库根目录> --知识库 <知识库文件夹> --输出 自检.xlsx
可选：--模板 <阅读笔记模板.md>（取它的二级标题作 4 级笔记的必有小节）；--笔记 <文件名的一部分>（只查这几篇，可重复）；
      --矩阵 <阅读矩阵.xlsx> 和 --取值表 <阅读矩阵取值表.json>（不给取值表时找矩阵旁边的同名文件）；
      --笔记目录 文献；--缓存 <目录>（PDF 文字缓存，重跑更快，放在库外）；--不核数字。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

FM = re.compile(r"\A---\s*\n(.*?)\n---\s*\n", re.S)
LINK = re.compile(r"(!?)\[\[([^\]|]+?)(?:\|([^\]]*))?\]\]")
PAGE_LINK = re.compile(r"\[\[[^\]|]+?\.pdf#page=\d+(?:\|[^\]]*)?\]\]", re.I)
ONE_PAGE = r"\[\[[^\]|]+?\.pdf#page=\d+(?:\|[^\]]*)?\]\]"
# 堆页码：句号后面跟着页码，或者不在括号里连着两个以上，或者一对括号里超过 3 处（连续页写成的范围算 1 处）
PAGE_ITEM = r"(?:" + ONE_PAGE + r")(?:—(?:" + ONE_PAGE + r"))?"
PILE = re.compile(r"[。！？]\s*(?:" + ONE_PAGE + r")|(?<![（(、—])(?:" + ONE_PAGE + r"){2,}|[（(](?:" + PAGE_ITEM + r"[、，,；;]?){4,}[）)]", re.I)
# 防御性限定的说法；“能不能”“是不是”这类问句和“而不是”这类对照不算
CAVEAT = re.compile(r"不能说明|不能证明|不能写成|不能把|不能将|不能据此|不能直接|不能用|不能当|不能算|不能替代|不能视为"
                    r"|不等于|不等同|不意味着|不代表|不授|不升级|不算|不作为|不足以|不宜|不应|不可|未必能")
INTERNAL = re.compile(r"原页|原图|图读|目视|表文冲突|口径|不授|不升级|不计新增|挂载|单人提取待核|如实保留|如实记为|不擅自|不偷偷|证据槽|槽位"
                      r"|擅自|冒充|补造|悄改|不猜|需核|仍待核|待核")
CODE = re.compile(r"(?<![A-Za-z0-9])(?:C[1-3]|E[1-5]|M[0-2]|D[0-3])(?![A-Za-z0-9])|(?<![A-Za-z])PF(?![A-Za-z])"
                  r"|本轮|外部补充|精读篇数|[CEM] ?等级|结构\d{3,4}")
CN_NUM = re.compile(r"(?<![第一二三四五六七八九十百千万两几])(?:[二三四五六七八九]?十[一二三四五六七八九]?|[三四五六七八九]|百)"
                    r"(?!分之)(?:个|篇|项|道|题|名|人|次(?![项方])|年|周|天|岁|波|组|类|页|张|条|所|种|份|处|档|轮|位|例"
                    r"|领域|维度|因子|时点|分量|侧面|章节|阶段|条目|指标|模型)")
BLOCK_GLUED = re.compile(r"(?<=\S)\^[A-Za-z0-9][\w-]*[ \t]*$", re.M)
LIMIT_TITLE = re.compile(r"边界|限度|不等于|不能|并非|而非")
# 把各段限定句挪出来集中存放的小节（reading-note-style.md 第 0 节第 7 条）；“补充分析”“补充材料”不算
ORPHAN_HEAD = re.compile(r"(?:方法|解释|说明|限定|注意).{0,3}补充|补充(?:说明|解释|限定)")
TEACH = re.compile(r"教学示例(（[^）]*）)?(?=\*{0,2}\s*[：:])")
# Obsidian 不显示 \( \) 和 \[ \]；\[ 只在里面像公式时才算，免得误报转义的方括号
BAD_MATH = re.compile(r"\\\(.+?\\\)|\\\[(?=[^\]\n]*(?:\\[A-Za-z]|[=^_]))")
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]|\r(?!\n)")
TOOLNAME = re.compile(r"[A-Za-z][A-Za-z]*(?:-[A-Za-z]+)*[- ]?\d+(?:\.\d+)?")
NUM = re.compile(r"(?<![A-Za-z0-9_.])(?:\d+\.\d+|\.\d+|\d{2,})(?![A-Za-z0-9_])")
NODE_TYPES = {"概念", "测量工具", "结局", "变量", "方法", "样本"}
FLOOR = {"测量工具": 2000, "方法": 2000, "概念": 1500, "样本": 1200, "结局": 1200, "变量": 1200}
RELATIONS = {"上位", "下位", "组成", "测量", "区分", "共变", "预测", "调节", "中介", "对照"}
D2_SECTIONS = ["为什么读这篇", "研究怎么做", "主要发现", "我的评价", "对我的文章有什么用"]
EVAL_H2 = {"我的评价", "对我的文章有什么用", "疑问与待核", "提取记录", "可能被追问的问题", "原文摘录"}
CODE_OK_H3 = {"证据等级为什么这样定"}
THEORY = re.compile(r"理论|综述|评论|述评|概念|方法学|方法论|社论")
SKIP_DIRS = {".obsidian", ".trash", ".git"}


def norm(s: str) -> str:
    return re.sub(r"\s+", "", s or "")


def frontmatter(text: str) -> tuple[dict[str, str], str]:
    m = FM.match(text)
    if not m:
        return {}, text
    meta = {}
    for line in m.group(1).splitlines():
        if ":" in line and not line[:1].isspace() and not line.startswith("-"):
            k, _, v = line.partition(":")
            meta[k.strip()] = v.strip()
    return meta, text[m.end():]


def links_in(value: str) -> list[str]:
    return [m.group(2).split("#")[0].split("/")[-1].strip() for m in LINK.finditer(value or "")]


def zi(text: str) -> int:
    """字数：汉字、英文词和数字各算 1，嵌入行、页码链接和注释不算。"""
    lines = [l for l in text.splitlines() if not re.match(r"^\s*>?\s*!\[\[", l)]
    s = re.sub(r"%%.*?%%", "", "\n".join(lines), flags=re.S)
    s = PAGE_LINK.sub("", s)
    s = LINK.sub(lambda m: m.group(3) or m.group(2).split("#")[0].split("/")[-1], s)
    return (len(re.findall(r"[一-鿿]", s)) + len(re.findall(r"[A-Za-z]+(?:[-'][A-Za-z]+)*", s))
            + len(re.findall(r"\d+(?:\.\d+)?", s)))


def blocks(body: str) -> list[tuple[str, str, str]]:
    """逐行返回 (所在二级标题, 所在三级标题, 这一行)。笔记里一段通常就是一行。"""
    out, h2, h3 = [], "", ""
    for line in body.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*?)\s*$", line)
        if m:
            if len(m.group(1)) == 2:
                h2, h3 = norm(m.group(2)), ""
            elif len(m.group(1)) == 3:
                h3 = norm(m.group(2))
            continue
        if line.strip():
            out.append((h2, h3, line))
    return out


def section(text: str, name: str) -> str:
    m = re.search(r"^## " + re.escape(name) + r"[^\n]*\n(.*?)(?=^## |\Z)", text, re.S | re.M)
    return m.group(1) if m else ""


class Index:
    def __init__(self, root: Path, skip: set[str]):
        self.by_rel: dict[str, Path] = {}
        self.by_name: dict[str, Path] = {}
        for p in root.rglob("*"):
            rel = p.relative_to(root)
            if any(part in skip for part in rel.parts) or not p.is_file():
                continue
            self.by_rel[rel.as_posix().lower()] = p
            self.by_name.setdefault(p.name.lower(), p)
            if p.suffix == ".md":
                self.by_name.setdefault(p.stem.lower(), p)
        self.cache: dict[Path, str] = {}

    def resolve(self, target: str) -> Path | None:
        low = target.strip().lower()
        for key in (low, low + ".md"):
            if key in self.by_rel:
                return self.by_rel[key]
        name = low.split("/")[-1]
        return self.by_name.get(name) or self.by_name.get(name + ".md")

    def text(self, p: Path) -> str:
        if p not in self.cache:
            self.cache[p] = p.read_text(encoding="utf-8", errors="replace")
        return self.cache[p]


class Pdfs:
    """PDF 文字，按页缓存。"""

    def __init__(self, cache: Path | None):
        self.cache = cache
        self.mem: dict[Path, list[str]] = {}
        if cache:
            cache.mkdir(parents=True, exist_ok=True)

    def pages(self, p: Path) -> list[str]:
        if p in self.mem:
            return self.mem[p]
        key = hashlib.md5(f"{p}|{p.stat().st_size}|{p.stat().st_mtime}".encode()).hexdigest()
        f = self.cache / f"{key}.json" if self.cache else None
        if f and f.exists():
            pages = json.loads(f.read_text(encoding="utf-8"))
        else:
            try:
                from pypdf import PdfReader
                pages = [(pg.extract_text() or "") for pg in PdfReader(str(p)).pages]
            except Exception:
                pages = []
            if f:
                f.write_text(json.dumps(pages, ensure_ascii=False), encoding="utf-8")
        self.mem[p] = [clean_pdf(t) for t in pages]
        return self.mem[p]


def clean_pdf(t: str) -> str:
    # 出版社字体把负号、等号存成代码：Wiley 的 /C0 是负号、/C6 是 ±；APA 的 /H11005 是等号，/H9251 是 α
    t = re.sub(r"/C0(?=[\d.])", "-", t).replace("/C6", "±")
    t = re.sub(r"/H11\d{3}|/H9\d{3}", " ", t)
    return re.sub(r"[−‒–—‐]", "-", t).replace("·", ".").replace("­", "")


def number_found(num: str, text: str) -> bool:
    """小数点也认逗号（欧洲期刊），整数也认千分位写法。"""
    s = num.rstrip("%")
    if s.startswith("0."):
        s = s[1:]
    if s.startswith("."):
        pats = [r"(?<![0-9])0?" + re.escape(s) + r"(?![0-9])", r"(?<![0-9])0?," + re.escape(s[1:]) + r"(?![0-9])"]
    elif "." in s:
        pats = [r"(?<![0-9.])" + re.escape(s) + r"(?![0-9])", r"(?<![0-9.,])" + re.escape(s.replace(".", ",")) + r"(?![0-9])"]
    else:
        pats = [r"(?<![0-9.,])" + re.escape(s) + r"(?![0-9])"]
        if len(s) > 3:
            pats.append(r"(?<![0-9.,])" + re.escape(f"{int(s):,}") + r"(?![0-9])")
    return any(re.search(p, text) for p in pats)


def check_links(text: str, where: str, idx: Index, pdfs: Pdfs, bad: list) -> None:
    for m in LINK.finditer(text):
        target, _, anchor = m.group(2).partition("#")
        if not target:
            continue
        p = idx.resolve(target)
        if p is None:
            bad.append((where, m.group(0)[:120], "找不到文件"))
            continue
        if not anchor:
            continue
        if anchor.startswith("page="):
            if p.suffix.lower() == ".pdf":
                n = len(pdfs.pages(p))
                page = int(re.sub(r"\D", "", anchor) or 0)
                if n and not 1 <= page <= n:
                    bad.append((where, m.group(0)[:120], f"PDF 只有 {n} 页"))
        elif anchor.startswith("^"):
            # Obsidian 只认前面有空格的块编号
            if not re.search(r"(?:^|\s)\^" + re.escape(anchor[1:]) + r"\s*$", idx.text(p), re.M):
                bad.append((where, m.group(0)[:120], "找不到块编号（或编号前缺空格）"))
        elif p.suffix == ".md":
            if not re.search(r"^#{1,6}\s+" + re.escape(anchor.strip()) + r"\s*$", idx.text(p), re.M):
                bad.append((where, m.group(0)[:120], "找不到小节标题"))


def check_numbers(note: str, h2: str, para: str, idx: Index, pdfs: Pdfs, nums: list, row: dict) -> None:
    pages: dict[Path, set[int]] = defaultdict(set)
    for m in re.finditer(r"\[\[([^\]|]+?\.pdf)#page=(\d+)", para, re.I):
        pdf = idx.resolve(m.group(1))
        if pdf:
            pages[pdf].add(int(m.group(2)))
    if not pages:
        return
    pool = []
    for pdf, ps in pages.items():
        allp = pdfs.pages(pdf)
        for n in ps:
            pool.extend(allp[i - 1] for i in (n - 1, n, n + 1) if 1 <= i <= len(allp))
    pool_text = "\n".join(pool)
    if len(pool_text.strip()) < 200:
        row["文字层"] = "无"
        return
    if len(re.findall(r"[A-Za-z0-9]", pool_text)) < 0.4 * len(re.sub(r"\s", "", pool_text)):
        row["文字层"] = "乱码"
        return
    if "虚构" in para:  # 教学示例里的数字是编的
        return
    plain = TOOLNAME.sub(" ", LINK.sub(" ", para))
    plain = re.sub(r"30 ?秒看懂|主张 ?\d+\.\d+-\d+-\d+|\d+\.\d+-\d+-\d+|结构\d{3,4}(?:[—–-]\d{3,4})?"
                   r"|95 ?%(?= ?(?:区间|置信|CI))|[pP] ?[<>≤≥] ?0?\.0?0?[15]\b", " ", plain)
    for m in NUM.finditer(plain):
        s = m.group(0)
        before = plain[:m.start()].rstrip()
        after = plain[m.end():].lstrip()[:1]
        if "." not in s and (int(s) < 30 or 1900 <= int(s) <= 2035):
            continue
        if after in ("页", "章", "节") or before.endswith(("第", "表", "图")):
            continue
        row["核数字"] = row.get("核数字", 0) + 1
        if number_found(s, pool_text):
            continue
        clause = re.split(r"[，。；,;]", plain[:m.start()])[-1] + re.split(r"[，。；,;]", plain[m.end():])[0]
        elsewhere = [i + 1 for pdf in pages for i, pg in enumerate(pdfs.pages(pdf)) if number_found(s, pg)]
        if re.search(r"算|相加|合计|换算|相减|减去|乘|除以|约", clause):
            kind = "可能是换算"
        elif elsewhere:
            kind = "页码可能不对"
        else:
            kind = "全文没找到"
        row["没找到"] = row.get("没找到", 0) + 1
        nums.append((note, h2, s, kind, "、".join(map(str, elsewhere[:4])),
                     plain[max(0, m.start() - 30):m.end() + 15].replace("\n", " "),
                     "；".join(f"{q.stem} 第 {','.join(map(str, sorted(v)))} 页" for q, v in pages.items())))


def check_note(p: Path, idx: Index, pdfs: Pdfs, d1_sections: list[str], verify: bool, bad: list, nums: list) -> dict:
    text = idx.text(p)
    meta, body = frontmatter(text)
    depth = meta.get("阅读深度", "") or meta.get("阅读层级", "")
    theory = bool(THEORY.search(meta.get("设计", "")))
    h2s = [norm(m.group(1)) for m in re.finditer(r"^## (.*?)\s*$", body, re.M)]
    h1 = next((m.group(1) for m in re.finditer(r"^# (.*?)\s*$", body, re.M)), "")
    head = "\n".join(body.splitlines()[:6])
    row = {"笔记": p.stem, "阅读深度": depth, "设计": meta.get("设计", ""), "状态": meta.get("状态", ""),
           "字数": zi(body), "理论类": "是" if theory else ""}
    issues = []
    if "[!summary]" not in head:
        issues.append("速览框不在开头")
    if LIMIT_TITLE.search(h1) or LIMIT_TITLE.search(p.stem):
        issues.append("标题写的是限制，应写发现或问题")
    if not theory:
        need = D2_SECTIONS if depth in ("D2", "3", "3 级") else d1_sections
        missing = [s for s in need if norm(s) not in h2s]
        if missing:
            issues.append("缺小节：" + "、".join(missing))
    floor = 6000 if depth in ("D1", "4", "4 级") else 2000 if depth in ("D2", "3", "3 级") else 0
    if floor and row["字数"] < floor:
        issues.append(f"字数 {row['字数']} 低于 {floor}")
    embeds = {m.group(2).split("#")[0].split("/")[-1].strip() for m in LINK.finditer(body)
              if m.group(1) and "#是什么" in m.group(2)}
    for key in ("测量工具", "方法"):
        lack = [n for n in links_in(meta.get(key, "")) if n not in embeds]
        if lack:
            issues.append(f"{key}没放补课框：" + "、".join(lack))
    glued = len(BLOCK_GLUED.findall(body))
    if glued:
        issues.append(f"{glued} 个块编号前缺空格")
    orphan = [h for h in (norm(m.group(1)) for m in re.finditer(r"^#{2,4}\s+(.*?)\s*$", body, re.M))
              if ORPHAN_HEAD.search(h) and h not in EVAL_H2]
    if orphan:
        issues.append("有收集限定句的小节：" + "、".join(orphan) + "（句子放回原段改写，删掉小节）")
    labels = [m.group(1) or "" for m in TEACH.finditer(body)]
    odd = [x for x in labels if x != "（虚构）"]
    if odd:
        issues.append(f"{len(odd)} 处教学示例标签不是“教学示例（虚构）”")
    math_bad, ctrl = len(BAD_MATH.findall(body)), len(CTRL.findall(text))
    if math_bad:
        issues.append(f"{math_bad} 处公式用了 \\( \\) 或 \\[ \\]，Obsidian 不显示，改用 $…$")
    if ctrl:
        issues.append(f"有 {ctrl} 个控制字符（多半是脚本把 \\f、\\a 这类转义吃掉了）")
    piles = caveat_paras = 0
    codes, cn, internal = Counter(), Counter(), Counter()
    for h2, h3, line in blocks(body):
        kind = line.lstrip()[:1]
        plain = LINK.sub("", line)
        piles += len(PILE.findall(line))
        in_eval = h2 in EVAL_H2 or h3 in CODE_OK_H3
        if kind not in (">", "|", "!") and not in_eval:
            if len(CAVEAT.findall(plain)) >= 2:
                caveat_paras += 1
            internal.update(m.group(0) for m in INTERNAL.finditer(plain))
        summary = kind == ">" and h2 == ""
        if not summary and not in_eval:
            codes.update(m.group(0) for m in CODE.finditer(plain))
        cn.update(m.group(0) for m in CN_NUM.finditer(plain))
        if verify and PAGE_LINK.search(line):
            check_numbers(p.stem, h2, line, idx, pdfs, nums, row)
    row.update({"句外或连着的页码": piles, "多处限定的段": caveat_paras,
                "讲解区代号": "、".join(f"{k}×{v}" for k, v in codes.most_common(6)),
                "讲解区内部用语": "、".join(f"{k}×{v}" for k, v in internal.most_common(6)),
                "中文计数": "、".join(f"{k}×{v}" for k, v in cn.most_common(5)),
                "补课框": len(re.findall(r"\[!tip\]", body)),
                "发现块": len(re.findall(r"\s\^f\d+\s*$", body, re.M)),
                "收集限定句的小节": "、".join(orphan), "教学示例标签不统一": len(odd), "公式问题": math_bad + ctrl,
                "问题": "；".join(issues)})
    check_links(body, p.stem, idx, pdfs, bad)
    return row


def check_card(p: Path, idx: Index, pdfs: Pdfs, uses: Counter, bad: list) -> dict:
    text = idx.text(p)
    meta, body = frontmatter(text)
    typ, tier = meta.get("类型", ""), meta.get("讲义", "")
    what = section(body, "是什么")
    n = zi(what)
    issues = []
    if not tier:
        issues.append("没定讲义档")
    if tier == "完整":
        if n < FLOOR.get(typ, 1200):
            issues.append(f"完整讲义 {n} 字，低于 {FLOOR.get(typ, 1200)}")
        if not re.search(r"^### ", what, re.M):
            issues.append("完整讲义没分小标题")
        if not section(body, "资料来源").strip(" \n%"):
            issues.append("没有资料来源")
    elif tier == "简短" and n < 300:
        issues.append(f"简短讲义只有 {n} 字")
    elif not tier and n < 300:
        issues.append(f"是什么只有 {n} 字")
    if uses[p.stem] >= 3 and tier != "完整":
        issues.append(f"已被 {uses[p.stem]} 篇笔记用到，应写完整讲义")
    lesson = LINK.sub("", what + section(body, "在本领域怎么用"))
    codes = Counter(m.group(0) for m in CODE.finditer(lesson))
    if codes:
        issues.append("讲义里有项目代号：" + "、".join(codes))
    if typ == "方法" and tier == "完整" and "$" not in body:
        issues.append("方法讲义没有公式（写出核心公式，用 $…$）")
    math_bad, ctrl = len(BAD_MATH.findall(body)), len(CTRL.findall(text))
    if math_bad:
        issues.append(f"{math_bad} 处公式用了 \\( \\) 或 \\[ \\]，Obsidian 不显示，改用 $…$")
    if ctrl:
        issues.append(f"有 {ctrl} 个控制字符（多半是脚本把 \\f、\\a 这类转义吃掉了）")
    for line in section(body, "关联").splitlines():
        if line.startswith("- "):
            parts = [x.strip() for x in line[2:].split("｜")]
            if len(parts) < 4 or parts[1] not in RELATIONS or parts[2] not in {"强", "中", "弱"}:
                issues.append("关联格式不对：" + line[:40])
                break
    check_links(body, p.stem, idx, pdfs, bad)
    return {"卡片": p.stem, "类型": typ, "讲义": tier, "是什么字数": n, "被几篇笔记用到": uses[p.stem],
            "公式问题": math_bad + ctrl, "问题": "；".join(issues)}


def check_matrix(path: Path, vocab: dict, notes: dict[str, dict], node_names: set[str]) -> list[tuple]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    ws = wb["数据"] if "数据" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    col = {norm(str(h or "")): i for i, h in enumerate(next(rows)) if h}
    id_col = vocab.get("编号列", "候选编号")
    state_col = vocab.get("状态列", "提取状态")
    done = set(vocab.get("已提取取值", ["已提取待核", "已核实"]))
    design_col, theory_value = vocab.get("设计列", "研究设计"), vocab.get("理论设计取值", "理论")
    controlled = {norm(k): set(v) for k, v in vocab.get("受控字段", {}).items()}
    prefixes = {norm(k): tuple(v) for k, v in vocab.get("开头取值", {}).items()}
    standard = set(vocab.get("标准方法名", []))
    placeholder = re.compile("|".join(vocab.get("占位词", ["待补", "TODO", "见笔记", "同上", r"\?\?", "？？"])))
    out, seen = [], set()
    for r in rows:
        get = lambda k: str(r[col[norm(k)]]).strip() if norm(k) in col and r[col[norm(k)]] is not None else ""
        cid = get(id_col)
        if get(state_col) not in done:
            continue
        seen.add(cid)
        theory = get(design_col) == theory_value
        for k, allowed in controlled.items():
            if k not in col:
                continue
            allowed = allowed | {"不适用"} if theory and vocab.get("理论文章可填不适用", True) else allowed
            v = get(k)
            if not v:
                out.append((cid, k, "空着", ""))
                continue
            wrong = [x for x in re.split(r"[；;、]", v) if x.strip() and x.strip() not in allowed]
            if wrong:
                out.append((cid, k, "不在取值表里", v[:80]))
        for k, pre in prefixes.items():
            v = get(k)
            if k in col and v and not v.startswith(pre):
                out.append((cid, k, "开头不是规定取值", v[:80]))
        for k in vocab.get("节点列", []):
            v = get(k)
            if v and v != "不适用":
                lack = [x.strip() for x in re.split(r"[；;]", v)
                        if x.strip() and x.strip() not in node_names and x.strip() not in standard]
                if lack:
                    out.append((cid, k, "名称对不上知识库节点或标准方法名", "、".join(lack)[:80]))
        for k, i in col.items():
            v = str(r[i] or "")
            if "[[" in v:
                out.append((cid, k, "单元格里有双链", v[:60]))
            elif placeholder.search(v):
                out.append((cid, k, "占位词", v[:60]))
        note = notes.get(cid)
        if not note:
            out.append((cid, "笔记", "矩阵已提取，库里找不到笔记", ""))
        elif note.get("阅读深度") and get("阅读深度") and note["阅读深度"] != get("阅读深度"):
            out.append((cid, "阅读深度", "与笔记不一致", f"矩阵 {get('阅读深度')}，笔记 {note['阅读深度']}"))
    for cid in notes:
        if cid and cid not in seen:
            out.append((cid, state_col, "有笔记但矩阵没标已提取", ""))
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--根", "--root", dest="root", required=True)
    ap.add_argument("--知识库", "--vault", dest="kb", required=True)
    ap.add_argument("--输出", "--output", dest="out", required=True)
    ap.add_argument("--模板", "--template", dest="template")
    ap.add_argument("--矩阵", "--matrix", dest="matrix")
    ap.add_argument("--取值表", "--vocabulary", dest="vocab")
    ap.add_argument("--笔记", "--note", dest="only", action="append", default=[])
    ap.add_argument("--笔记目录", "--notes-dir", dest="notes_dir", default="文献")
    ap.add_argument("--缓存", "--cache", dest="cache")
    ap.add_argument("--不核数字", "--no-numbers", dest="noverify", action="store_true")
    a = ap.parse_args()
    root = Path(a.root)
    kb = root / a.kb
    idx = Index(root, SKIP_DIRS)
    pdfs = Pdfs(Path(a.cache) if a.cache else None)
    d1 = D2_SECTIONS + ["研究问题与背景", "作者怎么解释", "关键图表怎么读"]
    if a.template:
        d1 = [norm(m.group(1)) for m in re.finditer(r"^## (.*?)\s*$", Path(a.template).read_text(encoding="utf-8"), re.M)]
    all_notes = sorted((kb / a.notes_dir).glob("*.md"))
    note_files = [p for p in all_notes if not a.only or any(o in p.stem for o in a.only)]
    card_files = [p for p in kb.rglob("*.md") if p.parent.name in NODE_TYPES]
    bad, nums = [], []
    note_rows = [check_note(p, idx, pdfs, d1, not a.noverify, bad, nums) for p in note_files]
    uses = Counter()
    for p in all_notes:
        uses.update({m.group(2).split("#")[0].split("/")[-1].strip() for m in LINK.finditer(idx.text(p))})
    card_rows = [] if a.only else [check_card(p, idx, pdfs, uses, bad) for p in card_files]
    matrix_rows = []
    if a.matrix and not a.only:
        vocab_path = Path(a.vocab) if a.vocab else Path(a.matrix).with_name("阅读矩阵取值表.json")
        vocab = json.loads(vocab_path.read_text(encoding="utf-8")) if vocab_path.exists() else {}
        if not vocab:
            print(f"没有找到取值表 {vocab_path}，矩阵只查双链、占位词和与笔记是否一致。")
        notes_meta = {frontmatter(idx.text(p))[0].get(vocab.get("编号列", "候选编号"), ""): frontmatter(idx.text(p))[0]
                      for p in all_notes}
        matrix_rows = check_matrix(Path(a.matrix), vocab, notes_meta, {p.stem for p in card_files})

    from openpyxl import Workbook
    wb = Workbook()
    summary = wb.active
    summary.title = "汇总"
    for title, rows in (("笔记", note_rows), ("知识卡", card_rows)):
        ws = wb.create_sheet(title)
        keys = list(dict.fromkeys(k for r in rows for k in r))
        ws.append(keys)
        for r in rows:
            ws.append([r.get(k, "") for k in keys])
    for title, head, rows in (("数字核对", ["笔记", "小节", "数字", "类别", "在全文哪几页找到", "上下文", "所引页"], nums),
                              ("坏链接", ["所在文件", "链接", "问题"], bad),
                              ("矩阵", ["编号", "字段", "问题", "内容"], matrix_rows)):
        ws = wb.create_sheet(title)
        ws.append(head)
        for r in rows:
            ws.append(list(r))
    lines = [
        ("笔记篇数", len(note_rows)),
        ("有版式问题的笔记", sum(1 for r in note_rows if r["问题"])),
        ("句外或连着的页码", sum(r["句外或连着的页码"] for r in note_rows)),
        ("多处限定的段", sum(r["多处限定的段"] for r in note_rows)),
        ("讲解区有代号的笔记", sum(1 for r in note_rows if r["讲解区代号"])),
        ("讲解区有内部用语的笔记", sum(1 for r in note_rows if r["讲解区内部用语"])),
        ("用中文数字计数的笔记", sum(1 for r in note_rows if r["中文计数"])),
        ("有收集限定句小节的笔记", sum(1 for r in note_rows if r["收集限定句的小节"])),
        ("教学示例标签不统一的笔记", sum(1 for r in note_rows if r["教学示例标签不统一"])),
        ("公式写法或控制字符有问题的笔记和知识卡", sum(1 for r in note_rows + card_rows if r["公式问题"])),
        ("核对的数字", sum(r.get("核数字", 0) for r in note_rows)),
        ("所引页上没找到的数字", len(nums)),
        ("  其中可能是换算", sum(1 for r in nums if r[3] == "可能是换算")),
        ("  其中页码可能不对", sum(1 for r in nums if r[3] == "页码可能不对")),
        ("  其中全文没找到", sum(1 for r in nums if r[3] == "全文没找到")),
        ("文字层缺失或乱码、没法核数字的笔记", sum(1 for r in note_rows if r.get("文字层"))),
        ("知识卡张数", len(card_rows)),
        ("有问题的知识卡", sum(1 for r in card_rows if r["问题"])),
        ("坏链接", len(bad)),
        ("矩阵问题", len(matrix_rows)),
    ]
    for k, v in lines:
        summary.append([k, v])
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    for k, v in lines:
        print(f"{k}：{v}")
    print(f"已写入 {out}")


if __name__ == "__main__":
    main()
